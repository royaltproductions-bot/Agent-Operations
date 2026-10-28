"""Bounded read-only document conversion for legitimate TryBounty attachments.

Convert common business files to plain text for the agent's existing work
context. No macros, external links, spreadsheet formulas, or active content are
executed. All callers must first cap incoming byte size.
"""
from __future__ import annotations

import csv
import io
from pathlib import Path
from zipfile import BadZipFile, ZipFile


TEXT_MIMES = {
    "text/plain", "text/markdown", "text/csv", "text/tab-separated-values",
    "application/json", "application/x-ndjson", "application/jsonl",
}
TEXT_EXTS = {".txt", ".md", ".csv", ".tsv", ".json", ".jsonl"}
DOCUMENT_MIMES = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
MAX_BINARY_BYTES = 3_000_000
MAX_DOCUMENT_CHARS = 24_000
MAX_OFFICE_UNCOMPRESSED_BYTES = 20_000_000


def kind_for(filename: str, mime: str) -> str:
    ext = Path(filename.lower()).suffix
    mime = mime.split(";")[0].strip().lower()
    if ext in DOCUMENT_MIMES:
        if mime in {"", "application/octet-stream", DOCUMENT_MIMES[ext]}:
            return ext
        raise ValueError("filename_mime_mismatch")
    if ext in TEXT_EXTS:
        if mime in TEXT_MIMES or mime in {"", "application/octet-stream"}:
            return "text"
        raise ValueError("filename_mime_mismatch")
    if mime in TEXT_MIMES and not ext:
        return "text"
    if mime in DOCUMENT_MIMES.values() and not ext:
        return next(ext for ext, typ in DOCUMENT_MIMES.items() if typ == mime)
    raise ValueError("unsupported_attachment_format")


def _check_office_archive(data: bytes):
    """Reject oversized decompression before loading user-supplied Office files."""
    try:
        with ZipFile(io.BytesIO(data)) as archive:
            parts = archive.infolist()
            if (len(parts) > 1500 or
                    sum(info.file_size for info in parts) > MAX_OFFICE_UNCOMPRESSED_BYTES):
                raise ValueError("office_archive_too_large")
    except BadZipFile as exc:
        raise ValueError("unreadable_office_archive") from exc


def convert_document(data: bytes, filename: str, mime: str,
                     limit: int = MAX_DOCUMENT_CHARS) -> str:
    """Extract only real, visible text; never claim to OCR a scanned PDF."""
    kind = kind_for(filename, mime)
    if kind == "text":
        text = data.decode("utf-8-sig")
    elif kind == ".pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data), strict=False)
            if reader.is_encrypted:
                raise ValueError("encrypted_document")
            pieces = []
            for number, page in enumerate(reader.pages):
                if number >= 60 or sum(map(len, pieces)) >= limit:
                    break
                extracted = (page.extract_text() or "").strip()
                if extracted:
                    pieces.append(f"[Page {number + 1}]\n{extracted}")
            text = "\n\n".join(pieces)
        except ImportError as exc:
            raise ValueError("pdf_parser_missing") from exc
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("unreadable_pdf") from exc
    elif kind == ".docx":
        try:
            from docx import Document
            _check_office_archive(data)
            document = Document(io.BytesIO(data))
            pieces = [p.text for p in document.paragraphs if p.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    pieces.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(pieces)
        except ImportError as exc:
            raise ValueError("docx_parser_missing") from exc
        except Exception as exc:
            raise ValueError("unreadable_docx") from exc
    else:  # .xlsx
        try:
            from openpyxl import load_workbook
            _check_office_archive(data)
            workbook = load_workbook(io.BytesIO(data), read_only=True,
                                     data_only=True, keep_links=False)
            pieces = []
            for sheet in workbook.worksheets[:8]:
                pieces.append(f"[Sheet: {sheet.title}]")
                out = io.StringIO(newline="")
                writer = csv.writer(out)
                for index, row in enumerate(sheet.iter_rows(values_only=True)):
                    if index >= 1000 or out.tell() >= limit:
                        break
                    writer.writerow([str(cell)[:300] if cell is not None else ""
                                     for cell in row[:40]])
                pieces.append(out.getvalue())
                if sum(map(len, pieces)) >= limit:
                    break
            workbook.close()
            text = "\n".join(pieces)
        except ImportError as exc:
            raise ValueError("xlsx_parser_missing") from exc
        except Exception as exc:
            raise ValueError("unreadable_xlsx") from exc
    text = text[:limit].strip()
    if not text:
        raise ValueError("no_extractable_text")
    return text


def csv_to_xlsx(data: bytes) -> bytes:
    """Render a validated, formula-neutralized CSV into an Office spreadsheet."""
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Results"
    for row in csv.reader(io.StringIO(data.decode("utf-8"), newline="")):
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        letter = column[0].column_letter
        width = min(50, max(12, max(len(str(cell.value or ""))
                                    for cell in column[:200]) + 2))
        sheet.column_dimensions[letter].width = width
    output = io.BytesIO()
    workbook.save(output)
    if output.tell() > MAX_BINARY_BYTES:
        raise ValueError("xlsx_deliverable_too_large")
    return output.getvalue()
