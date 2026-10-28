"""Real file round-trips for the document types Gambit now handles."""
import io
from importlib.util import find_spec
import unittest
from zipfile import ZIP_DEFLATED, ZipFile

from tools import gambit_documents as documents


class DocumentTests(unittest.TestCase):
    def test_format_checks_do_not_trust_mismatched_content_types(self):
        self.assertEqual(documents.kind_for("leads.xlsx", "application/octet-stream"), ".xlsx")
        with self.assertRaises(ValueError):
            documents.kind_for("leads.pdf", "image/png")
        with self.assertRaises(ValueError):
            documents.kind_for("clip.mp4", "video/mp4")

    def test_plain_text_conversion(self):
        self.assertEqual(documents.convert_document(b'Name,City\nAcme,Augusta',
                                                     'leads.csv', 'text/csv'),
                         'Name,City\nAcme,Augusta')

    def test_small_compressed_office_bomb_is_rejected(self):
        binary = io.BytesIO()
        with ZipFile(binary, 'w', compression=ZIP_DEFLATED) as archive:
            archive.writestr('word/document.xml', 'x' * 21_000_000)
        self.assertLess(len(binary.getvalue()), documents.MAX_BINARY_BYTES)
        with self.assertRaises(ValueError):
            documents.convert_document(binary.getvalue(), 'brief.docx',
                                       documents.DOCUMENT_MIMES['.docx'])

    @unittest.skipUnless(find_spec('docx'), 'python-docx not installed')
    def test_word_document_paragraph_and_table_are_read(self):
        from docx import Document
        document = Document()
        document.add_paragraph('Business research brief')
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Acme'
        table.cell(0, 1).text = 'Augusta'
        data = io.BytesIO()
        document.save(data)
        text = documents.convert_document(data.getvalue(), 'brief.docx',
                                          documents.DOCUMENT_MIMES['.docx'])
        self.assertIn('Business research brief', text)
        self.assertIn('Acme | Augusta', text)

    @unittest.skipUnless(find_spec('openpyxl'), 'openpyxl not installed')
    def test_excel_inputs_and_outputs_are_real_workbooks(self):
        from openpyxl import Workbook, load_workbook
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(['Company', 'City'])
        worksheet.append(['Acme', 'Augusta'])
        data = io.BytesIO()
        workbook.save(data)
        text = documents.convert_document(data.getvalue(), 'leads.xlsx',
                                          documents.DOCUMENT_MIMES['.xlsx'])
        self.assertIn('Acme,Augusta', text)

        from tools.gambit_worker import prepare_csv
        rendered = documents.csv_to_xlsx(prepare_csv(
            'company,score\nAcme,5\n=2+3,7'))
        loaded = load_workbook(io.BytesIO(rendered), read_only=True, data_only=False)
        self.assertEqual(loaded.active['A2'].value, 'Acme')
        self.assertEqual(loaded.active['A3'].value, "'=2+3")
        self.assertEqual(loaded.active['A3'].data_type, 's')

    @unittest.skipUnless(find_spec('pypdf'), 'pypdf not installed')
    def test_pdf_extracts_text_instead_of_accepting_a_filename(self):
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
        writer = PdfWriter()
        page = writer.add_blank_page(width=250, height=200)
        font = writer._add_object(DictionaryObject({
            NameObject('/Type'): NameObject('/Font'),
            NameObject('/Subtype'): NameObject('/Type1'),
            NameObject('/BaseFont'): NameObject('/Helvetica')}))
        page[NameObject('/Resources')] = DictionaryObject({
            NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
        content = DecodedStreamObject()
        content.set_data(b'BT /F1 12 Tf 30 170 Td (Acme Corp) Tj ET')
        page[NameObject('/Contents')] = writer._add_object(content)
        data = io.BytesIO()
        writer.write(data)
        text = documents.convert_document(data.getvalue(), 'company.pdf',
                                          documents.DOCUMENT_MIMES['.pdf'])
        self.assertIn('Acme Corp', text)


if __name__ == '__main__':
    unittest.main()
